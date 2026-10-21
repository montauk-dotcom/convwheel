"""Создаёт тестовые файлы разных форматов в tests/fixtures."""
import subprocess, sys
from pathlib import Path
import docx, openpyxl
from PIL import Image, ImageDraw
from pptx import Presentation
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from convwheel.platform_utils import find_ffmpeg

D = Path(__file__).parent / "fixtures"; D.mkdir(exist_ok=True)
d = docx.Document(); d.add_heading("Отчёт по проекту", 1)
d.add_paragraph("Привет, мир! Это проверка кириллицы и «кавычек».")
t = d.add_table(rows=2, cols=3)
for i, row in enumerate([["Имя","Сумма","Дата"],["Алия","1 200","21.10.2026"]]):
    for j, v in enumerate(row): t.cell(i,j).text = v
d.add_paragraph("Последний абзац."); d.save(D/"Отчёт.docx")
(D/"заметка_cp1251.txt").write_bytes("Строка в кодировке Windows-1251\nВторая строка".encode("cp1251"))
(D/"readme.md").write_text("# Заголовок\n\nТекст с **жирным** и [ссылкой](http://x.y).\n\n- пункт 1\n- пункт 2\n", encoding="utf-8")
wb = openpyxl.Workbook(); ws = wb.active; ws.title="Продажи"; ws.append(["Товар","Кол-во"]); ws.append(["Чай",5]); ws.append(["Итого","=B2*2"])
ws2 = wb.create_sheet("Склад"); ws2.append(["A",1]); wb.save(D/"таблица.xlsx")
(D/"data.csv").write_text("name;qty;price\nЯблоко;3;1,5\nГруша;007;2\n", encoding="utf-8")
p = Presentation(); s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text="Слайд 1"; s.placeholders[1].text="Пункт"; p.save(D/"deck.pptx")
im = Image.new("RGBA",(300,200),(0,0,0,0)); ImageDraw.Draw(im).ellipse((20,20,280,180),fill=(40,120,255,255)); im.save(D/"logo.png")
frames=[Image.new("RGB",(80,80),(i*40,100,200)) for i in range(5)]; frames[0].save(D/"anim.gif",save_all=True,append_images=frames[1:],duration=120,loop=0)
ff = find_ffmpeg()
subprocess.run([ff,"-y","-loglevel","error","-f","lavfi","-i","testsrc=size=321x241:rate=25:duration=3","-f","lavfi","-i","sine=frequency=440:duration=3","-shortest","-c:v","mpeg4","-c:a","aac",str(D/"clip.avi")],check=True)
subprocess.run([ff,"-y","-loglevel","error","-f","lavfi","-i","sine=frequency=440:duration=2",str(D/"tone.wav")],check=True)
print("ok", sorted(x.name for x in D.iterdir()))
