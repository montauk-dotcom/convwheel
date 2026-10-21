"""Прогоняет ВСЕ варианты конвертации для каждого тестового файла и проверяет содержимое."""
import shutil, sys, time
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from convwheel.formats import targets_for, targets_for_many, MERGE_PDF
from convwheel.engine import run_job
from convwheel.converters.common import ConversionError

SRC = Path(__file__).parent / "fixtures"
WORK = Path(__file__).parent / "out"; shutil.rmtree(WORK, ignore_errors=True); WORK.mkdir()

def check(out: Path, key: str):
    import pymupdf, docx
    if out.is_dir():
        files = list(out.iterdir()); assert files, "пустая папка"; return f"папка, {len(files)} файлов"
    ext = out.suffix.lower()[1:]
    if ext == "pdf":
        d = pymupdf.open(out); txt = "".join(p.get_text() for p in d); return f"{d.page_count} стр, текст: {txt[:40]!r}"
    if ext == "docx":
        return "текст: " + repr(" ".join(p.text for p in docx.Document(out).paragraphs)[:50])
    if ext in ("txt","csv","html","md","rtf"):
        t = out.read_text(encoding="utf-8-sig"); return repr(t[:50])
    if ext in ("png","jpg","webp","ico","gif","bmp","tiff"):
        from PIL import Image; im = Image.open(out); return f"{im.format} {im.size} {im.mode} frames={getattr(im,'n_frames',1)}"
    if ext in ("xlsx",):
        import openpyxl; ws = openpyxl.load_workbook(out).active; return repr([[c.value for c in r] for r in ws.iter_rows()][:3])
    return f"{out.stat().st_size} байт"

fails = 0
for f in sorted(SRC.iterdir()):
    local = WORK / f.name; shutil.copy(f, local)
    for t in targets_for(local):
        t0 = time.time()
        try:
            outs = run_job([local], t)
            print(f"OK  {f.name:22} → {t.key:5} {time.time()-t0:4.1f}s  {outs[0].name}: {check(outs[0], t.key)}")
        except ConversionError as e:
            fails += 1; print(f"ERR {f.name:22} → {t.key:5} {e}")
# пакетный режим: несколько картинок → один PDF
imgs = [WORK/"logo.png", WORK/"anim.gif"]
print("many:", [t.key for t in targets_for_many(imgs)])
print("merge:", check(run_job(imgs, MERGE_PDF)[0], "pdf"))
# повторная конвертация не перезаписывает
o1 = run_job([WORK/"logo.png"], [t for t in targets_for(WORK/"logo.png") if t.key=="jpg"][0])[0]
print("unique name:", o1.name)
print("FAILS:", fails)
