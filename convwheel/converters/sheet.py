"""Таблицы без офисного пакета (запасной путь): XLSX ↔ CSV через openpyxl."""
from __future__ import annotations

import csv
import re
from pathlib import Path

from .common import ConversionError, Ctx, ensure_output, read_text_smart


def xlsx_to_csv(src: Path, ctx: Ctx) -> list[Path]:
    try:
        import openpyxl
    except ImportError as e:
        raise ConversionError("Нужен openpyxl или LibreOffice") from e
    wb = openpyxl.load_workbook(str(src), read_only=True, data_only=True)
    outs = []
    sheets = wb.worksheets
    for ws in sheets:
        suffix = "" if len(sheets) == 1 else "-" + re.sub(r'[\\/:*?"<>|]', "_", ws.title)
        out = ctx.out(src, "csv", suffix)
        # utf-8-sig — чтобы Excel на Windows корректно открыл кириллицу двойным кликом
        with open(out, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                w.writerow(["" if v is None else v for v in row])
        outs.append(out)
    wb.close()
    return ensure_output(outs)


def csv_to_xlsx(src: Path, ctx: Ctx) -> list[Path]:
    import openpyxl
    text = read_text_smart(src)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = src.stem[:31] or "Sheet1"
    for row in csv.reader(text.splitlines(), dialect):
        ws.append([_num(v) for v in row])
    out = ctx.out(src, "xlsx")
    wb.save(str(out))
    return ensure_output([out])


def _num(v: str):
    s = v.strip()
    if re.fullmatch(r"-?\d+", s) and not (len(s) > 1 and s.lstrip("-").startswith("0")) and len(s) < 16:
        return int(s)
    if re.fullmatch(r"-?\d+[.,]\d+", s):
        try:
            return float(s.replace(",", "."))
        except ValueError:
            pass
    return v
