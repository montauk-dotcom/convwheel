"""Конвертации текста без офисного пакета: DOCX→TXT, TXT/MD→DOCX, MD→HTML."""
from __future__ import annotations

import html
import re
from pathlib import Path

from .common import ConversionError, Ctx, ensure_output, read_text_smart


def docx_to_txt(src: Path, ctx: Ctx) -> list[Path]:
    """Текст из DOCX в порядке документа: абзацы и таблицы (ячейки через табуляцию)."""
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as e:
        raise ConversionError("Нужен пакет python-docx") from e
    try:
        d = docx.Document(str(src))
    except Exception as e:
        raise ConversionError(f"Не удалось открыть DOCX: {e}") from e

    lines: list[str] = []
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(Paragraph(child, d).text)
        elif tag == "tbl":
            table = Table(child, d)
            for row in table.rows:
                cells, prev = [], None
                for c in row.cells:          # объединённые ячейки повторяются — убираем дубли
                    if c._tc is prev:
                        continue
                    prev = c._tc
                    cells.append(c.text.replace("\n", " ").strip())
                lines.append("\t".join(cells))
            lines.append("")
    text = "\n".join(lines).strip() + "\n"
    out = ctx.out(src, "txt")
    out.write_text(text, encoding="utf-8")
    return ensure_output([out])


def text_to_docx(src: Path, ctx: Ctx) -> list[Path]:
    import docx
    from docx.shared import Pt

    text = read_text_smart(src)
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    is_md = src.suffix.lower() == ".md"
    for line in text.splitlines():
        if is_md:
            m = re.match(r"^(#{1,6})\s+(.*)", line)
            if m:
                d.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
                continue
            m = re.match(r"^\s*[-*+]\s+(.*)", line)
            if m:
                d.add_paragraph(_strip_md(m.group(1)), style="List Bullet")
                continue
            m = re.match(r"^\s*\d+[.)]\s+(.*)", line)
            if m:
                d.add_paragraph(_strip_md(m.group(1)), style="List Number")
                continue
            line = _strip_md(line)
        d.add_paragraph(line)
    out = ctx.out(src, "docx")
    d.save(str(out))
    return ensure_output([out])


def _strip_md(s: str) -> str:
    s = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), s)
    s = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*|_(.+?)_", lambda m: m.group(1) or m.group(2), s)
    s = re.sub(r"`([^`]+)`", r"\1", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    return s


def markdown_to_html(md_text: str, title: str = "") -> str:
    try:
        import markdown
        body = markdown.markdown(md_text, extensions=["tables", "fenced_code", "sane_lists"])
    except ImportError:
        body = f"<pre>{html.escape(md_text)}</pre>"
    return (
        "<!doctype html><html><head><meta charset='utf-8'>"
        f"<title>{html.escape(title)}</title>"
        "<style>body{font-family:-apple-system,Segoe UI,Arial,sans-serif;max-width:820px;"
        "margin:40px auto;line-height:1.55;padding:0 16px}pre,code{background:#f3f3f3}"
        "table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:4px 8px}</style>"
        f"</head><body>{body}</body></html>"
    )


def text_to_html(src: Path, ctx: Ctx) -> list[Path]:
    text = read_text_smart(src)
    if src.suffix.lower() == ".md":
        page = markdown_to_html(text, src.stem)
    else:
        page = markdown_to_html("", src.stem).replace(
            "<body></body>", f"<body><pre style='white-space:pre-wrap'>{html.escape(text)}</pre></body>")
    out = ctx.out(src, "html")
    out.write_text(page, encoding="utf-8")
    return ensure_output([out])


def text_to_txt(src: Path, ctx: Ctx) -> list[Path]:
    """MD → TXT (без разметки) или перекодировка TXT в UTF-8."""
    text = read_text_smart(src)
    if src.suffix.lower() == ".md":
        text = "\n".join(_strip_md(re.sub(r"^#{1,6}\s+", "", ln)) for ln in text.splitlines())
    out = ctx.out(src, "txt")
    out.write_text(text, encoding="utf-8")
    return ensure_output([out])
